# -*- coding: utf-8 -*-
# Рисовалка для build3.py: терминалы, графики, карта доклада, мини-карта, схемы узлов — всё нативными фигурами.
# Подключается через exec() после «головы» build.py и build2.py (оттуда clone/text/card/zones/row/ctext/...).
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION

INK, MUTED, FAINT = '1E5B28', '6B7280', 'C9CED6'
TERM_BG, TERM_FG = '10241A', 'E6F2E9'
FRAMES = os.path.join(HERE, 'frames')


def rect(s, x, y, w, h, fill, line=None, radius=0.12, lw=1.0, dash=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / max(min(w, h), 0.01))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
        if dash: sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def seg(s, x1, y1, x2, y2, color=MUTED, w=1.25, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(w)
    if dash: c.line.dash_style = MSO_LINE.DASH
    return c


def box(s, x, y, w, h, label, fill='FFFFFF', line=None, color=INK, size=11, mono=False, bold=False, radius=0.1, align='c'):
    rect(s, x, y, w, h, fill, line, radius=radius)
    text(s, x + 0.06, y, w - 0.12, h, label if not mono else [('`%s`' % l) for l in ([label] if isinstance(label, str) else label)],
         size=size, color=color, align=align, anchor='m', font='X5 Sans Medium' if bold else 'X5 Sans')


def term(s, x, y, w, h, lines, size=13, caption=None, hl=()):
    """Терминал с реальным выводом. lines — список строк или имя файла из frames/."""
    if isinstance(lines, str):
        lines = open(os.path.join(FRAMES, lines), encoding='utf-8').read().rstrip('\n').split('\n')
    rect(s, x, y, w, h, TERM_BG, radius=0.16)
    for i, c in enumerate(('FF6B6B', 'FFD166', 'AADC00')):
        rect(s, x + 0.2 + i * 0.2, y + 0.16, 0.11, 0.11, c, shape=MSO_SHAPE.OVAL)
    if caption:
        text(s, x + 1.0, y + 0.1, w - 1.2, 0.25, caption, size=9, color='8FB09A', align='r')
    tb = s.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.42), Inches(w - 0.5), Inches(h - 0.55)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    rx = re.compile('(' + '|'.join(re.escape(k) for k in hl) + ')') if hl else None
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        base = 'AADC00' if ln.startswith('$') else ('8FB09A' if ln.startswith('#') or ln.startswith('(') else TERM_FG)
        parts = rx.split(ln) if (rx and not ln.startswith('$')) else [ln]
        for part in parts:
            if part == '': continue
            r = p.add_run(); r.text = part; f = r.font; f.name = MONO; f.size = Pt(size)
            f.color.rgb = RGBColor.from_string('FF9DB8' if (hl and part in hl) else base)
            f.bold = bool(hl and part in hl)


def code(s, x, y, w, h, lines, size=13, accent=None, hl=(), hlcolor=None, label=None):
    """Светлая карточка кода; строки из hl подсвечиваются."""
    clone(s, 'white' if accent != 'red' else 'pinkcard', x, y, w, h)
    top = y + 0.2
    if label:
        lw = 0.1 * len(label) + 0.4
        clone(s, 'lime' if accent != 'red' else 'pink', x + 0.2, y + 0.16, lw, 0.34)
        text(s, x + 0.2, y + 0.16, lw, 0.34, label, size=11, color=INK if accent != 'red' else BRED, align='c', anchor='m', font='X5 Sans Medium')
        top = y + 0.6
    tb = s.shapes.add_textbox(Inches(x + 0.3), Inches(top), Inches(w - 0.5), Inches(h - (top - y) - 0.15)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(2)
        r = p.add_run(); r.text = ln if ln else ' '; f = r.font; f.name = MONO; f.size = Pt(size)
        on = any(k in ln for k in hl)
        f.color.rgb = RGBColor.from_string((hlcolor or (BRED if accent == 'red' else INK)) if on else '6B7280'); f.bold = on


# ---------- графики ----------
def _style_chart(ch, size=11):
    ch.font.size = Pt(size); ch.font.name = 'X5 Sans'; ch.font.color.rgb = RGBColor.from_string(INK)
    ch.has_legend = False; ch.has_title = False


def bar_chart(s, x, y, w, h, cats, vals, colors, fmt='0.00', size=11, vmax=None, labels=True):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series('v', vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd); ch = gf.chart
    _style_chart(ch, size)
    pl = ch.plots[0]; pl.gap_width = 55; pl.vary_by_categories = False
    ser = pl.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = RGBColor.from_string(c)
    if labels:
        pl.has_data_labels = True; dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(size + 3); dl.font.name = 'X5 Sans Medium'; dl.font.color.rgb = RGBColor.from_string(INK)
    from pptx.enum.chart import XL_TICK_LABEL_POSITION, XL_TICK_MARK
    va = ch.value_axis; va.has_major_gridlines = False; va.minimum_scale = 0
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE; va.major_tick_mark = XL_TICK_MARK.NONE; va.format.line.fill.background()
    if vmax: va.maximum_scale = vmax
    ca = ch.category_axis; ca.format.line.color.rgb = RGBColor.from_string(FAINT); ca.tick_labels.font.size = Pt(size)
    return ch


def line_chart(s, x, y, w, h, cats, series, fmt='"×"0.00', size=11, vmax=None):
    cd = CategoryChartData(); cd.categories = cats
    for name, vals, _ in series: cd.add_series(name, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd); ch = gf.chart
    _style_chart(ch, size)
    for ser, (_, _, col) in zip(ch.plots[0].series, series):
        ser.format.line.color.rgb = RGBColor.from_string(col); ser.format.line.width = Pt(3); ser.smooth = False
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(col); ser.marker.size = 9
        ser.marker.format.line.color.rgb = RGBColor.from_string(col)
    last = ch.plots[0].series[-1]
    last.data_labels.show_value = True; last.data_labels.number_format = fmt; last.data_labels.number_format_is_linked = False
    last.data_labels.position = XL_LABEL_POSITION.ABOVE; last.data_labels.font.size = Pt(size + 3); last.data_labels.font.name = 'X5 Sans Medium'
    last.data_labels.font.color.rgb = RGBColor.from_string(BRED)
    va = ch.value_axis; va.minimum_scale = 0; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = RGBColor.from_string('E5E7EB')
    va.format.line.fill.background(); va.tick_labels.font.size = Pt(size - 1); va.tick_labels.number_format = '"×"0'; va.tick_labels.number_format_is_linked = False
    if vmax: va.maximum_scale = vmax
    ch.category_axis.format.line.color.rgb = RGBColor.from_string(FAINT)
    return ch


# ---------- карта доклада: 12 шагов, две дорожки ----------
STEPS = [  # (название, когда, A, B, чем увидеть)
    ('две функции', 'хук', '0.12 мс', '0.79 мс', 'shape-analyzer'),
    ('форма объекта', 'вызов 1', '1 форма', '8 форм', '%HaveSameMap'),
    ('место чтения', 'первые вызовы', 'MONO', 'MEGA', '--log-ic'),
    ('ярусы', 'первые миллисекунды', 'одна лестница', 'одна лестница', '--max-opt'),
    ('ставка', 'после компиляции', 'сыграла', 'wrong map → общий путь', '--trace-deopt'),
    ('сервис', 'прод', '2302 rps', '1180 rps', 'autocannon'),
    ('агент', 'до первого вызова', 'автор: ?', 'без правил · 30/30', 'claude -p + прибор'),
    ('корпус', 'до первого вызова', 'живое', 'мёртвое рядом', 'матрица по версиям'),
    ('правила', 'до первого вызова', 'с файлом · 30/30', 'без файла', 'v8-rules.md'),
    ('лев', 'верхняя ступень', 'ставка та же', 'ставка та же', '--turbolev'),
    ('карта целиком', '—', 'форма + файл', 'форма + нет файла', 'как проверить у себя'),
    ('снова слайд 1', 'финал', '1 → MONO → 2302', '8 → MEGA → 1180', 'репо'),
]
BREATHERS = (3, 5, 8)


def talk_map(s, x, y, w, h, upto=12, current=None, size=None):
    """Одна линия из 12 шагов; пройденное проявлено, текущее обведено, дальнейшее приглушено. Масштабируется по h."""
    k = max(0.8, min(1.7, h / 3.0)); fs = (size or 8) * (0.85 + 0.25 * k)
    n = len(STEPS); lab_w = 0.5 * k; step = (w - lab_w) / n
    ax = y + 0.62 * k; ra = ax + 1.0 * k; rb = ax + 1.55 * k; rt = ax + 2.12 * k; d = 0.36 * k
    seg(s, x + lab_w, ax, x + w - 0.05, ax, FAINT, 2)
    for nm, ry, col in (('A', ra, INK), ('B', rb, BRED)):
        rect(s, x, ry, 0.34 * k, 0.4 * k, 'E0FBCC' if nm == 'A' else 'FFD9E4', radius=0.08)
        text(s, x, ry, 0.34 * k, 0.4 * k, nm, size=fs + 2, color=col, align='c', anchor='m', font='X5 Sans Medium')
    for i, (name, when, a_, b_, tool) in enumerate(STEPS, 1):
        cx = x + lab_w + step * (i - 0.5); on = i <= upto; cur = (i == current)
        if cur: rect(s, cx - d * 0.72, ax - d * 0.72, d * 1.44, d * 1.44, 'AADC00', shape=MSO_SHAPE.OVAL)
        rect(s, cx - d / 2, ax - d / 2, d, d, INK if on else 'E5E7EB', shape=MSO_SHAPE.OVAL)
        text(s, cx - d / 2, ax - d / 2, d, d, str(i), size=fs + 2, color='FFFFFF' if on else '9CA3AF', align='c', anchor='m', font='X5 Sans Medium')
        text(s, cx - step / 2, y, step, 0.3 * k, when, size=fs - 1.5, color=MUTED if on else FAINT, align='c', anchor='m')
        text(s, cx - step / 2 + 0.02, ax + 0.3 * k, step - 0.04, 0.6 * k, name, size=fs + 1.5, color=INK if on else 'B6BCC6', align='c', font='X5 Sans Medium')
        if on:
            for val, ry, col, bg in ((a_, ra, INK, 'F1FBE8'), (b_, rb, BRED, 'FFEAF0')):
                rect(s, cx - step / 2 + 0.03, ry, step - 0.06, 0.44 * k, bg, radius=0.06)
                text(s, cx - step / 2 + 0.05, ry, step - 0.1, 0.44 * k, val, size=fs - 0.5, color=col, align='c', anchor='m')
            text(s, cx - step / 2 + 0.02, rt, step - 0.04, 0.4 * k, tool, size=fs - 1.5, color=MUTED, align='c')
        if i in BREATHERS:
            rect(s, cx + step / 2 - 0.09 * k, ax - 0.09 * k, 0.18 * k, 0.18 * k, 'FFD166' if i <= upto else 'E5E7EB', shape=MSO_SHAPE.DIAMOND)


def minimap(s, x, y, w, h, current):
    clone(s, 'white', x, y, w, h)
    text(s, x + 0.15, y + 0.1, w - 0.3, 0.22, 'карта доклада', size=9, color=MUTED)
    n = 12; st = (w - 0.4) / (n - 1); cy = y + 0.62
    seg(s, x + 0.2, cy, x + w - 0.2, cy, FAINT, 1.5)
    for i in range(1, n + 1):
        cx = x + 0.2 + st * (i - 1); d = 0.26 if i == current else 0.13
        fill = 'AADC00' if i == current else (INK if i < current else 'D9DDE3')
        rect(s, cx - d / 2, cy - d / 2, d, d, fill, line=INK if i == current else None, lw=1.5, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.12, y + 0.9, w - 0.24, h - 1.0, [f'**шаг {current} из 12**', STEPS[current - 1][0]], size=11, color=INK, align='c', anchor='m', gap=2)


# ---------- схемы узлов ----------
def schema_tree(s, x, y, w, h):
    """Узел 2: A — одна карточка формы; B — три развилки, восемь листьев."""
    aw = 2.7
    clone(s, 'white', x, y, aw, h)
    text(s, x + 0.2, y + 0.1, aw - 0.4, 0.35, '**A** · `?? null`', size=15)
    fields = ['id', 'name', 'email', 'plan', 'credits', 'referrer', 'trialUntil', 'teamId']
    rh = (h - 1.15) / 8
    for i, f in enumerate(fields):
        rect(s, x + 0.25, y + 0.5 + i * rh, aw - 0.5, rh - 0.05, 'F1FBE8' if i < 5 else 'E0FBCC', radius=0.05)
        text(s, x + 0.35, y + 0.5 + i * rh, aw - 0.7, rh - 0.05, f'`{f}`' + ('  ?? null' if i >= 5 else ''), size=13, anchor='m')
    text(s, x, y + h - 0.65, aw, 0.55, '**1 форма**', size=22, align='c', anchor='m')
    bx = x + aw + 0.25; bw = w - aw - 0.25
    clone(s, 'pinkcard', bx, y, bw, h)
    text(s, bx + 0.2, y + 0.1, bw - 0.4, 0.35, '**B** · условный спред', size=15, color=BRED)
    levels = [('5 полей', 1), ('referrer?', 2), ('trialUntil?', 4), ('teamId?', 8)]
    ly = [y + 0.55 + k * ((h - 1.5) / 3) for k in range(4)]
    pos = {}
    for k, (lab, cnt) in enumerate(levels):
        for j in range(cnt if k > 0 else 1):
            pass
    # узлы решений: уровень k имеет 2^k узлов; листья — 8
    names = ['referrer?', 'trialUntil?', 'teamId?']
    for k in range(3):
        cnt = 2 ** k; bwid = min(1.35, (bw - 0.5) / cnt - 0.08)
        for j in range(cnt):
            cx = bx + 0.25 + (bw - 0.5) * (j + 0.5) / cnt
            pos[(k, j)] = (cx, ly[k])
    leaves = ['—', '+i', '+t', '+t +i', '+r', '+r +i', '+r +t', '+r +t +i']
    for j in range(8):
        pos[(3, j)] = (bx + 0.25 + (bw - 0.5) * (j + 0.5) / 8, ly[3])
    for k in range(3):
        for j in range(2 ** k):
            px, py = pos[(k, j)]
            for c in (2 * j, 2 * j + 1):
                qx, qy = pos[(k + 1, c)]
                seg(s, px, py + 0.24, qx, qy - 0.06, 'E6A3B7', 1.25)
    for k in range(3):
        cnt = 2 ** k; bwid = min(1.7, (bw - 0.5) / cnt - 0.1)
        for j in range(cnt):
            cx, cy = pos[(k, j)]
            box(s, cx - bwid / 2, cy - 0.22, bwid, 0.46, names[k], fill='E3F5F9', size=11 if k == 2 else 13, mono=True)
    lw = (bw - 0.5) / 8 - 0.06
    for j in range(8):
        cx, cy = pos[(3, j)]
        box(s, cx - lw / 2, cy - 0.06, lw, 0.46, leaves[j], fill='FFD9E4', color=BRED, size=10)
    text(s, bx, y + h - 0.6, bw, 0.5, '**2³ = 8 форм**', size=22, color=BRED, align='c', anchor='m')


def schema_ic(s, x, y, w, h):
    """Узел 3: место чтения u.plan — MONO / POLY / MEGA."""
    cw = (w - 0.4) / 3
    spec = (('MONO', '1 форма', 1, 'быстрый путь: одна проверка, одно чтение', 'lime', INK, 'читатель объектов A'),
            ('POLY', '2–4 формы', 3, 'перебор короткого списка форм', 'blue', INK, ''),
            ('MEGA', '5+ форм', 8, 'общий словарь: поиск по имени поля и форме', 'pink', BRED, 'читатель объектов B'))
    for i, (nm, cnt, k, cap, key, col, who) in enumerate(spec):
        cx = x + i * (cw + 0.2)
        clone(s, key, cx, y, cw, h)
        text(s, cx + 0.2, y + 0.12, cw - 0.4, 0.45, f'**{nm}**', size=26, color=col)
        text(s, cx + 0.2, y + 0.2, cw - 0.4, 0.3, cnt, size=14, color=col, align='r')
        cs = min(0.26, (cw - 0.5) / 8 - 0.03)
        for j in range(k):
            rect(s, cx + 0.22 + j * (cs + 0.04), y + 0.7, cs, cs, 'FFFFFF', line=col, radius=0.04)
            text(s, cx + 0.22 + j * (cs + 0.04), y + 0.7, cs, cs, str(j + 1), size=8, color=col, align='c', anchor='m')
        seg(s, cx + 0.22 + cs / 2, y + 0.7 + cs, cx + 0.22 + cs / 2, y + 1.22, col, 1.25)
        box(s, cx + 0.22, y + 1.22, 1.3, 0.44, 'u.plan', fill='FFFFFF', line=col, color=col, size=14, mono=True)
        text(s, cx + 0.22, y + 1.8, cw - 0.44, h - 2.3, cap, size=15, color=col)
        if who: text(s, cx + 0.22, y + h - 0.5, cw - 0.44, 0.35, who, size=13, color=col, font='X5 Sans Medium')


TIERS = (('Ignition', 'интерпретатор', 'байткод, без предположений', 505), ('Sparkplug', 'быстрый компилятор', 'без предположений', 775),
         ('Maglev', 'компилятор со ставкой', 'смотрит записи мест чтения', 70), ('TurboFan', 'лучший компилятор', 'со ставкой', 38))


def schema_tiers(s, x, y, w, h, top_label=None, numbers=True):
    """Узлы 4 и 10: лестница по числу вызовов."""
    base = y + h - 0.75; sw = (w - 0.2) / 4; hs = (0.9, 1.35, 2.05, 2.75)
    scale = (h - 1.75) / 2.75
    seg(s, x, base, x + w, base, INK, 1.75)
    text(s, x + w - 2.6, base + 0.4, 2.6, 0.25, 'количество вызовов функции →', size=11, color=MUTED, align='r')
    ticks = ('вызов 1', 'десятки', 'сотни', 'тысячи')
    for i, (nm, role, note, ms) in enumerate(TIERS):
        bx = x + 0.1 + i * sw; bh = hs[i] * scale
        spec = i >= 2
        fill = 'E0FBCC' if spec else 'F1F3F5'
        if top_label and i == 3: fill = 'AADC00'
        rect(s, bx, base - bh, sw - 0.06, bh, fill, line=INK if (top_label and i == 3) else None, lw=2, radius=0.08)
        lab = [f'**{top_label if (top_label and i == 3) else nm}**', role]
        if numbers: lab.append(f'**{ms} мс**')
        text(s, bx + 0.08, base - bh + 0.04, sw - 0.22, bh - 0.08, lab, size=11 if bh < 1.0 else 13, align='c', anchor='m', gap=1)
        seg(s, bx, base, bx, base + 0.1, INK, 1.25)
        text(s, bx - 0.02, base + 0.12, sw, 0.3, ticks[i], size=12, color=INK, font='X5 Sans Medium')
    # скобки
    by = y + 0.05
    for (i0, i1, lab, col) in ((0, 1, 'читают записи мест чтения, ничего не предполагают', MUTED), (2, 3, 'вшивают ставку «сюда приходит одна форма» — проверка одним сравнением', INK)):
        x0 = x + 0.1 + i0 * sw; x1 = x + 0.1 + (i1 + 1) * sw - 0.06
        seg(s, x0, by + 0.42, x1, by + 0.42, col, 1.25); seg(s, x0, by + 0.42, x0, by + 0.52, col, 1.25); seg(s, x1, by + 0.42, x1, by + 0.52, col, 1.25)
        text(s, x0, by - 0.05, x1 - x0, 0.45, lab, size=11.5, color=col, align='c', anchor='m')
    text(s, x + 0.1, base + 0.4, w - 3.0, 0.25, 'движок считает не вызовы, а исполненный байткод: функция с циклом внутри — TurboFan к ~1500-му вызову, крошечная без циклов — к ~10–15 тыс. · по времени — первые 2–5 мс работы', size=9.5, color=MUTED)


def chips(s, x, y, labels, h=0.5, fill='FFFFFF', line=INK, color=INK, size=12.5, gap=0.3, mark=None):
    cx = x
    for i, lab in enumerate(labels):
        wd = 0.098 * len(lab) + 0.32
        box(s, cx, y, wd, h, lab, fill=fill, line=line, color=color, size=size)
        if i < len(labels) - 1:
            seg(s, cx + wd, y + h / 2, cx + wd + gap, y + h / 2, line, 1.25)
            if mark is not None and i == mark:
                rect(s, cx + wd + gap / 2 - 0.07, y + h / 2 - 0.07, 0.14, 0.14, BRED, shape=MSO_SHAPE.OVAL)
        cx += wd + gap
    return cx


def schema_deopt(s, x, y, w, h):
    """Узел 5: ставка сыграла / wrong map."""
    lh = (h - 0.2) / 2
    clone(s, 'white', x, y, w, lh)
    text(s, x + 0.25, y + 0.12, w - 0.5, 0.3, '**A** · читатель объектов A на оптимизирующем ярусе', size=14)
    text(s, x + 0.25, y + 0.12, w - 0.5, 0.3, 'записи: MONO, 1 форма', size=11.5, color=MUTED, align='r')
    ex = chips(s, x + 0.3, y + 0.7, ['проверка формы', 'быстрый путь вшит'])
    seg(s, ex - 0.3, y + 0.95, x + w - 0.3, y + 0.95, '3FA34D', 3)
    text(s, x + 0.3, y + 1.3, w - 0.6, 0.4, 'ставка играет каждый вызов · у читателя строки wrong map нет (0 из 5 прогонов)', size=12, color=MUTED)
    by = y + lh + 0.2
    clone(s, 'pinkcard', x, by, w, lh)
    text(s, x + 0.25, by + 0.12, w - 0.5, 0.3, '**B** · читатель объектов B на оптимизирующем ярусе', size=14, color=BRED)
    text(s, x + 0.25, by + 0.12, w - 0.5, 0.3, 'сначала одна форма → потом восемь', size=11.5, color=MUTED, align='r')
    ex = chips(s, x + 0.3, by + 0.7, ['ставка на одну форму', 'откат', 'записи обновлены', 'перекомпиляция без ставки'], line=BRED, color=BRED, mark=0)
    seg(s, ex - 0.3, by + 0.95, x + w - 0.3, by + 0.95, 'D6336C', 3)
    text(s, x + 2.9, by + 0.36, 3.4, 0.3, '↑ `reason: wrong map`', size=11, color=BRED)
    text(s, x + 0.3, by + 1.3, w - 0.6, 0.5, 'один откат (5 из 5 прогонов) · дальше — общий путь MEGA, стационарно · в приборе ставки не было вовсе: 0 откатов, TurboFan, и всё равно ×6.6', size=11.5, color=MUTED)


def qr_png(url, path):
    import qrcode
    q = qrcode.QRCode(border=1, box_size=12); q.add_data(url); q.make(fit=True)
    q.make_image(fill_color='#1E5B28', back_color='white').save(path)
    return path
