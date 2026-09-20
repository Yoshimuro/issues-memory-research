# -*- coding: utf-8 -*-
import copy, os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

SRC = '/Users/User/Projects/Personal/autumn26holyjs/06-talk-final.md'
TPL = os.environ.get('X5_TEMPLATE', os.path.expanduser('~/Downloads/А. Зайцев.pptx'))
OUT = sys.argv[1]

GREEN, LIME, PALE, RED, GREY, WHITE = '1E5B28', 'AADC00', 'E0FBCC', 'C8102E', '6B7280', 'FFFFFF'
MONO = 'Menlo'

# ---------- сценарий ----------
md = open(SRC, encoding='utf-8').read()
script = md.split('# Послайдовый сценарий')[1].split('## Порядок первых появлений')[0]
SL, NODES = {}, {}
for m in re.finditer(r'^## Узел (\d+)\. (.+?) — ([\d.]+) мин\n(.*?)(?=^## Узел |\Z)', script, re.S | re.M):
    n, title, mins, body = int(m[1]), m[2], m[3], m[4]
    q = re.search(r'> \*\*Вопрос зала после узла:\*\* (.+)', body)
    br = re.search(r'> \*\*Передышка\.\*\* (.+)', body)
    NODES[n] = dict(title=title, mins=mins, q=q[1].strip() if q else '', breather=br[1].strip() if br else '')
    for s in re.finditer(r'^### (\d+\.\d+) · (.+?) — ([\d.]+) мин\n(.*?)(?=^### |^> \*\*|\Z)', body, re.S | re.M):
        f = {}
        for k in ('На экране', 'Визуал', 'Шапки A/B', 'Карта', 'Демо', 'Фолбэк', 'Речь'):
            fm = re.search(r'\*\*' + re.escape(k) + r'\.\*\*\s*(.*?)(?=\n\*\*[А-ЯA-Z][^*]*[.:]\*\*|\Z)', s[4], re.S)
            f[k] = fm[1].strip() if fm else ''
        SL[s[1]] = dict(title=s[2], mins=s[3], **f)

def ab_state(sid):
    raw = SL[sid]['Шапки A/B']
    if not raw or 'нет' in raw[:16]:
        return None
    raw = re.split(r' — (?=[а-яё])| \(', raw)[0]
    a, b = [], []
    for item in raw.split(' · '):
        k, _, v = item.strip().partition(' ')
        va, _, vb = v.partition('/')
        a.append(f'{k} {va or "—"}'); b.append(f'{k} {(vb or va) or "—"}')
    return ' · '.join(a), ' · '.join(b)

# ---------- шаблон ----------
prs = Presentation(TPL)
T = list(prs.slides)
def proto(si, shape_id):
    for sh in T[si - 1].shapes:
        if sh.shape_id == shape_id:
            return sh._element
    raise KeyError((si, shape_id))
P = dict(white=proto(2, 4), blue=proto(2, 13), dark=proto(10, 8), lime=proto(7, 7), shadow=proto(15, 3),
         pinkcard=proto(18, 4), pink=proto(19, 13), pill=proto(4, 9), darkbox=proto(4, 14),
         chev=proto(32, 7), badge=proto(21, 32))
LAY = dict(title=T[0].slide_layout, green=T[1].slide_layout, white=T[2].slide_layout,
           grey=T[9].slide_layout, dark=T[3].slide_layout)
N_TPL = len(T)
_id = [1000]

def clone(s, key, x, y, w, h):
    el = copy.deepcopy(P[key])
    _id[0] += 1
    el.find(qn('p:nvSpPr')).find(qn('p:cNvPr')).set('id', str(_id[0]))
    xf = el.find(qn('p:spPr')).find(qn('a:xfrm'))
    xf.find(qn('a:off')).set('x', str(Inches(x))); xf.find(qn('a:off')).set('y', str(Inches(y)))
    xf.find(qn('a:ext')).set('cx', str(Inches(w))); xf.find(qn('a:ext')).set('cy', str(Inches(h)))
    tb = el.find(qn('p:txBody'))
    if tb is not None:
        for p in tb.findall(qn('a:p'))[1:]:
            tb.remove(p)
        p0 = tb.find(qn('a:p'))
        for r in list(p0):
            if r.tag != qn('a:pPr'):
                p0.remove(r)
    s.shapes._spTree.append(el)
    return el

def text(s, x, y, w, h, paras, size=20, color=GREEN, align='l', anchor='t', font='X5 Sans', bold=False, gap=0, ls=None):
    """paras: строка или список строк; `код` → моноширинный; **x** → X5 Sans Medium."""
    if isinstance(paras, str):
        paras = [paras]
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = dict(t=MSO_ANCHOR.TOP, m=MSO_ANCHOR.MIDDLE, b=MSO_ANCHOR.BOTTOM)[anchor]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = dict(l=PP_ALIGN.LEFT, c=PP_ALIGN.CENTER, r=PP_ALIGN.RIGHT)[align]
        if gap:
            p.space_after = Pt(gap)
        if ls:
            p.line_spacing = ls
        for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', para):
            if not part:
                continue
            r = p.add_run()
            f = r.font
            f.size = Pt(size); f.color.rgb = RGBColor.from_string(color); f.name = font; f.bold = bold
            if part.startswith('`'):
                r.text = part[1:-1]; f.name = MONO; f.size = Pt(size * 0.88)
            elif part.startswith('**'):
                r.text = part[2:-2]; f.name = 'X5 Sans Medium'
            else:
                r.text = part
    return tb

def todo(s, x, y, w, h, label, kind='АССЕТ'):
    """Серая пунктирная заглушка под ассет, которого ещё нет."""
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.08, 0.25 / max(h, 0.1) * 0.5)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string('F2F2F2')
    sh.line.color.rgb = RGBColor.from_string('9CA3AF'); sh.line.width = Pt(1.5); sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    pad = 0.25
    text(s, x + pad, y + pad * 0.6, w - 2 * pad, h - pad * 1.2, [f'**{kind}**', label], size=14, color=GREY, align='c', anchor='m', gap=4)
    return sh

def notes(s, sid):
    d = SL[sid]
    parts = [f'[{sid} · {d["mins"]} мин]', '', d['Речь'], '', '— — —']
    for k in ('Визуал', 'Карта', 'Демо', 'Фолбэк'):
        if d[k]:
            parts.append(f'{k.upper()}: {d[k]}')
    s.notes_slide.notes_text_frame.text = '\n'.join(parts)

def strip(s, sid):
    st = ab_state(sid)
    if not st:
        return
    for key, x, lab, val, col in (('lime', 0.37, 'A', st[0], GREEN), ('pink', 6.75, 'B', st[1], '8A1538')):
        clone(s, key, x, 6.78, 6.2, 0.42)
        text(s, x + 0.2, 6.78, 5.8, 0.42, f'**{lab}** · {val}', size=9.5, color=col, anchor='m')

def set_title(s, t, color=None):
    s.shapes.title.text_frame.text = t
    if color:
        for r in s.shapes.title.text_frame.paragraphs[0].runs:
            r.font.color.rgb = RGBColor.from_string(color)

def drop_empty_placeholders(s):
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1 and not ph.text_frame.text:
            ph._element.getparent().remove(ph._element)

def light(sid, bg='green', title=None, title_w=None):
    s = prs.slides.add_slide(LAY[bg])
    d = SL[sid]
    set_title(s, title or d['title'])
    if title_w:
        tt = s.shapes.title
        tt.left, tt.top, tt.width, tt.height = Inches(0.37), Inches(0.40), Inches(title_w), Inches(0.58)
    drop_empty_placeholders(s)
    text(s, 11.4, 0.42, 1.55, 0.3, sid, size=11, color=GREY, align='r')
    strip(s, sid); notes(s, sid)
    return s

def dark(pill, big, box, sid=None, note=None, big_size=None):
    s = prs.slides.add_slide(LAY['dark'])
    set_title(s, big)
    tt = s.shapes.title
    tt.left, tt.top, tt.width, tt.height = Inches(1.18), Inches(2.0), Inches(10.98), Inches(2.1)
    for p in tt.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        if big_size:
            for r in p.runs:
                r.font.size = Pt(big_size)
    drop_empty_placeholders(s)
    pw = max(2.6, 0.155 * len(pill) + 0.8)
    clone(s, 'pill', (13.333 - pw) / 2, 1.2, pw, 0.5)
    text(s, (13.333 - pw) / 2, 1.2, pw, 0.5, pill, size=15, color=GREEN, align='c', anchor='m', font='X5 Sans Medium')
    if box:
        clone(s, 'darkbox', 1.99, 4.55, 9.36, 1.5)
        text(s, 2.3, 4.6, 8.74, 1.4, box, size=19, color=PALE, align='c', anchor='m')
    if sid:
        notes(s, sid)
    elif note:
        s.notes_slide.notes_text_frame.text = note
    return s

def divider(n):
    nd = NODES[n]
    dark(f'Узел {n} из 12 · {nd["mins"]} мин', nd['title'], ['Вопрос зала:', NODES[n - 1]['q']],
         note=f'Разделитель узла {n}. Вопрос, которым закончился узел {n-1}, — им же открывается этот.\n\n{NODES[n-1]["q"]}', big_size=40)

def breather(n, k):
    s = prs.slides.add_slide(LAY['white'])
    set_title(s, f'Передышка {k} · карта, 5 секунд без текста')
    drop_empty_placeholders(s)
    todo(s, 0.47, 1.3, 12.38, 5.6, 'lifecycle-map · состояние после узла %d — без текста и без речи' % n, 'СХЕМА')
    s.notes_slide.notes_text_frame.text = 'ПЕРЕДЫШКА. ' + NODES[n]['breather']

def card(s, key, x, y, w, h, paras=None, size=18, color=GREEN, align='l', anchor='m', pad=0.3, **kw):
    clone(s, key, x, y, w, h)
    if paras:
        text(s, x + pad, y + pad * 0.6, w - 2 * pad, h - pad * 1.2, paras, size=size, color=color, align=align, anchor=anchor, **kw)

def row(s, key, items, y, h, x0=0.47, x1=12.85, gapx=0.3, **kw):
    n = len(items); w = (x1 - x0 - gapx * (n - 1)) / n
    for i, it in enumerate(items):
        card(s, key, x0 + i * (w + gapx), y, w, h, it, **kw)

def ab_cards(s, y, h, a, b, size=16, labels=('A', 'B'), **kw):
    """Сплит A (зелёная рамка) / B (красная рамка)."""
    for key, x, lab, paras, col in (('white', 0.47, labels[0], a, GREEN), ('pinkcard', 6.82, labels[1], b, '8A1538')):
        clone(s, key, x, y, 6.03, h)
        clone(s, 'lime' if key == 'white' else 'pink', x + 0.25, y + 0.2, max(0.6, 0.13 * len(lab) + 0.4), 0.42)
        text(s, x + 0.25, y + 0.2, max(0.6, 0.13 * len(lab) + 0.4), 0.42, lab, size=13, color=col, align='c', anchor='m', font='X5 Sans Medium')
        text(s, x + 0.35, y + 0.75, 5.35, h - 0.95, paras, size=size, color=GREEN, anchor='m', **kw)

def bullets(s, x, y, w, items, size=20, step=1.15):
    for i, it in enumerate(items):
        clone(s, 'chev', x, y + i * step + 0.05, 0.2, 0.55)
        text(s, x + 0.5, y + i * step, w - 0.5, step - 0.15, it, size=size, anchor='m')

def big(s, x, y, w, h, t, size=72, color=GREEN, align='c'):
    text(s, x, y, w, h, t, size=size, color=color, align=align, anchor='m', font='X5 Sans Medium')

# ================= СЛАЙДЫ =================
# 1.1 титул
s = prs.slides.add_slide(LAY['title'])
set_title(s, 'V8 для самых маленьких')
for ph in s.placeholders:
    if ph.placeholder_format.type != 1:
        ph.text_frame.text = 'Что за лев этот турболев?'
text(s, 0.37, 4.35, 6.2, 0.6, 'В описании доклада часть текста написал агент. Честно.', size=13, color=PALE)
text(s, 0.37, 5.21, 5.5, 0.9, ['**Александр Зайцев**', 'HolyJS 2026 Autumn · 24.10.2026'], size=16, color=WHITE, gap=2)
notes(s, '1.1')

# 1.2
s = light('1.2', 'white')
dim = ['`id, name, email, plan, credits` — 5 строк одинаковы', '']
ab_cards(s, 1.3, 3.9, dim + ['`referrer: raw.referrer ?? null`', '`trialUntil: raw.trialUntil ?? null`', '`teamId: raw.teamId ?? null`'],
         dim + ['`...(raw.referrer && { referrer: raw.referrer })`', '`...(raw.trialUntil && { trialUntil: … })`', '`...(raw.teamId && { teamId: … })`'], size=17, gap=6)
row(s, 'blue', ['чексуммы равны', 'тесты зелёные', 'ревью: approve'], 5.4, 1.0, size=18, align='c')

# 1.3
s = light('1.3', 'green')
ab_cards(s, 1.3, 2.6, ['**0.37 мс**'], ['**1.58 мс**'], size=60, align='c')
todo(s, 0.47, 4.1, 8.2, 2.4, 'R1 · прогон shape-analyzer на ref-a.js и ref-b.js: shapes / ic / p50 / correct', 'ЗАПИСЬ')
text(s, 9.0, 4.1, 3.85, 2.4, ['**~×4.3**', '`lab/shape-analyzer.js`', 'Node 24.18.0 / V8 13.6', 'p50 горячего цикла «собрать и прочитать»'], size=14, color=GREY, gap=6, anchor='m')

# 1.4
s = light('1.4', 'white')
card(s, 'pinkcard', 0.47, 1.3, 8.4, 1.25, '`[bailout … reason: wrong map]`', size=24, align='c')
card(s, 'dark', 9.15, 1.3, 3.7, 1.25, 'улика, не причина', size=20, color=WHITE, align='c')
todo(s, 0.47, 2.8, 12.38, 3.7, 'lifecycle-map · состояние 1: всё серое, светятся «файл A» и «файл B», «автор: ?»', 'СХЕМА')

# узел 2
divider(2)
s = light('2.1', 'grey')
clone(s, 'dark', 0.37, 1.34, 4.74, 5.1)
text(s, 0.69, 1.7, 4.2, 1.2, 'форма', size=54, color=WHITE, font='X5 Sans Medium')
text(s, 0.69, 5.3, 4.2, 0.9, ['hidden class', 'в трейсе — `map`'], size=15, color=PALE)
bullets(s, 5.59, 1.5, 7.3, ['внутреннее описание объекта: какие поля, в каком порядке добавлены', 'одинаково собранные объекты делят одно описание'], size=22, step=1.5)
card(s, 'white', 5.59, 4.55, 3.5, 1.9, ['**A**: 8 полей', 'форм: **1**'], size=20, align='c')
card(s, 'pinkcard', 9.35, 4.55, 3.5, 1.9, ['**B**', 'форм: **?**'], size=20, align='c')

s = light('2.2', 'white')
card(s, 'white', 0.47, 1.3, 3.6, 5.2, ['**A**', 'одна карточка формы', '', '**1**'], size=22, align='c')
todo(s, 4.35, 1.3, 5.6, 5.2, 'дерево: «5 полей» → referrer ± → trialUntil ± → teamId ± → 8 листьев (прообраз — 03-storyboard, схема 2)', 'СХЕМА')
clone(s, 'pinkcard', 10.2, 1.3, 2.65, 5.2)
big(s, 10.2, 2.4, 2.65, 1.6, '2³ = 8', size=44, color='8A1538')
text(s, 10.4, 4.2, 2.25, 1.2, 'форм у B', size=18, align='c', color='8A1538')

s = light('2.3', 'green')
todo(s, 0.47, 1.3, 8.4, 5.2, 'LIVE №1 · `node --allow-natives-syntax lab/demo-samemap.js` → true / false / «форм у B: 8». Фолбэк-скрин R3 — следующим кадром', 'LIVE-ДЕМО')
card(s, 'dark', 9.15, 1.3, 3.7, 2.5, ['`%HaveSameMap`', 'A → true', 'B → false'], size=20, color=WHITE, gap=4)
clone(s, 'white', 9.15, 4.05, 3.7, 2.45)
big(s, 9.15, 4.05, 3.7, 2.45, '1 vs 8', size=48)

s = light('2.4', 'white')
todo(s, 0.47, 1.3, 12.38, 5.2, 'lifecycle-map · состояние 2: клетка «сборка / форма» — A 1, B 8; зелёная дорожка и красный пучок из восьми', 'СХЕМА')

# узел 3
divider(3)
s = light('3.1', 'green')
ab_cards(s, 1.3, 3.6, ['сборка', '**×0.92–1.02**'], ['чтение', '**×4.1–5.8**'], size=44, align='c', labels=('normalize', 'consume'))
todo(s, 0.47, 5.15, 8.4, 1.35, 'две пары столбцов B/A из lab/real/agg.json', 'ГРАФИК')
text(s, 9.15, 5.2, 3.7, 1.3, ['lab/real · Node 24', 'пара агента той же формы, 1 vs 8 — чья, дойдём'], size=13, color=GREY, anchor='m', gap=4)

s = light('3.2', 'grey')
clone(s, 'dark', 0.37, 1.34, 4.74, 5.1)
text(s, 0.69, 1.7, 4.2, 2.6, ['место чтения', 'запоминает формы'], size=34, color=WHITE, font='X5 Sans Medium')
text(s, 0.69, 5.5, 4.2, 0.6, 'inline cache', size=16, color=PALE)
for i, (k, v, key) in enumerate((('MONO', '1 форма — быстрый путь', 'lime'), ('POLY', '2–4 формы — перебор', 'blue'), ('MEGA', '5+ форм — общий словарь', 'pink'))):
    card(s, key, 5.59, 1.34 + i * 1.2, 7.26, 1.05, f'**{k}** · {v}', size=21)
clone(s, 'white', 5.59, 5.0, 7.26, 1.44)
big(s, 5.59, 5.0, 7.26, 1.44, 'MEGA vs MONO ×3.2–4.1', size=32)

s = light('3.3', 'white')
todo(s, 0.47, 1.3, 8.4, 3.6, 'R4 · 3–5 строк LoadIC по полю u.plan для B (1 → P → N) + одна строка для A; колонка состояния подсвечена', 'СКРИН')
card(s, 'dark', 9.15, 1.3, 3.7, 3.6, ['**1 → P → N**', '', 'A: одна строка, MONO'], size=24, color=WHITE, align='c')
text(s, 0.47, 5.0, 12.38, 0.4, '`node --log-ic --no-logfile-per-isolate --logfile=ic.log`', size=14, color=GREY)
todo(s, 0.47, 5.5, 12.38, 1.0, 'карта, состояние 3: «место чтения» A MONO / B MEGA; на «сборке» ×0.92–1.02', 'СХЕМА')

s = dark('Цена объяснена. Улика — ещё нет', 'Два вопроса', ['1. Спасёт ли компилятор B?', '2. Растёт ли цена с числом чтений?'], sid='3.4')
breather(3, 1)

# узел 4
divider(4)
s = light('4.1', 'green', title='Спасёт ли компилятор B?')
clone(s, 'dark', 0.47, 1.3, 3.9, 5.2)
big(s, 0.47, 1.3, 3.9, 5.2, 'Нет.', size=80, color=WHITE)
card(s, 'white', 4.65, 1.3, 3.95, 3.0, ['**A**', '**0.37 мс**'], size=36, align='c')
card(s, 'pinkcard', 8.9, 1.3, 3.95, 3.0, ['**B**', '**1.58 мс**'], size=36, align='c')
card(s, 'blue', 4.65, 4.55, 8.2, 0.9, 'горячий цикл = давно скомпилированный код', size=19, align='c')
text(s, 4.65, 5.7, 8.2, 0.7, 'у компилятора есть кое-что своё → по вызовам', size=16, color=GREY, align='c', anchor='m')

s = light('4.2', 'white', title='Ярус: четыре способа исполнить одну функцию')
for i, (tier, calls) in enumerate((('Ignition', '1'), ('Sparkplug', '≈ 8'), ('Maglev', '≈ 400'), ('TurboFan', '≈ 10 000'))):
    x = 0.47 + i * 3.17
    card(s, 'lime', x, 3.6 - i * 0.7, 2.87, 1.3 + i * 0.7, [f'**{tier}**', f'вызов {calls}'], size=20, align='c')
text(s, 0.47, 5.05, 8.0, 0.4, '«≈» — бюджет по длине байткода · Node 24 / V8 13.6', size=14, color=GREY)
todo(s, 0.47, 5.5, 12.38, 1.0, 'R5 · 3 строки `--print-bytecode` normalizeUser (CreateObjectLiteral / GetNamedProperty / Return)', 'СКРИН')

s = light('4.3', 'green')
clone(s, 'white', 0.47, 1.3, 12.38, 5.2)
big(s, 0.9, 1.5, 11.5, 1.5, '1417 → 999 → 61 → 41 мс', size=54)
row(s, 'blue', ['`--max-opt=0 / 1 / 2 / 3`', '**×34**', 'одна функция, не A против B'], 3.3, 1.4, x0=0.94, x1=12.4, size=19, align='c')
todo(s, 0.94, 4.95, 7.5, 1.3, 'лестница Ignition → Sparkplug → Maglev → TurboFan, ось вызовов 1 / ≈8 / ≈400 / ≈10 000', 'ГРАФИК')
text(s, 8.7, 4.95, 3.7, 1.3, 'подъём A и B по отдельности — нет замера', size=14, color=GREY, anchor='m')

s = light('4.4', 'grey')
card(s, 'shadow', 0.64, 1.4, 5.8, 4.9, ['**Ignition · Sparkplug**', '', 'читают записи мест чтения,', 'ничего не предполагают'], size=24)
card(s, 'dark', 6.89, 1.4, 5.8, 4.9, ['**Maglev · TurboFan**', '', 'вшивают в машинный код ставку', '«сюда приходит одна форма»', '', 'проверка — одно сравнение'], size=24, color=WHITE)

# узел 5
divider(5)
s = light('5.1', 'white', title='Ставка не сыграла: wrong map')
card(s, 'white', 0.47, 1.3, 4.2, 5.2, ['**A**', '', 'проверка формы ✓', '→ быстрый путь вшит'], size=22, align='c')
steps = ['проверка формы ✗', '`bailout … reason: wrong map`', 'откат в интерпретатор', 'записи обновлены', 'перекомпиляция']
for i, st in enumerate(steps):
    card(s, 'pinkcard' if i != 1 else 'pink', 4.95, 1.3 + i * 0.9, 5.4, 0.78, st, size=17, align='c', pad=0.15)
card(s, 'dark', 10.6, 1.3, 2.25, 4.38, ['термин', '**деопт**'], size=24, color=WHITE, align='c')
todo(s, 4.95, 5.85, 7.9, 0.65, 'R2 · строка из --trace-deopt на B (только после проверки 1!)', 'СКРИН')

s = light('5.2', 'green', title='Стационар: общий путь')
clone(s, 'white', 0.47, 1.3, 12.38, 5.2)
big(s, 0.9, 1.5, 11.5, 1.7, '1.58 = цена общего пути', size=56)
row(s, 'blue', ['после нескольких откатов компилятор перестаёт ставить', 'общий путь MEGA — на каждом чтении', 'деопт — событие, не налог'], 3.45, 1.7, x0=0.94, x1=12.4, size=17, align='c')
text(s, 0.94, 5.4, 11.46, 0.8, 'число деоптов — нет замера', size=14, color=GREY, align='c', anchor='m')

s = light('5.3', 'grey', title='Улика прочитана')
card(s, 'pinkcard', 0.47, 1.3, 6.0, 2.4, '`reason: wrong map`', size=30, align='c')
card(s, 'lime', 0.47, 3.95, 2.85, 2.55, ['**map**', '= форма'], size=24, align='c')
card(s, 'lime', 3.62, 3.95, 2.85, 2.55, ['**wrong**', '= ставка'], size=24, align='c')
clone(s, 'dark', 6.8, 1.3, 6.05, 5.2)
big(s, 6.8, 1.5, 6.05, 2.2, '2008', size=96, color=WHITE)
text(s, 7.2, 3.9, 5.25, 2.3, ['формы и записи мест чтения — с первого релиза V8', '', 'старше всех ступеней лестницы'], size=20, color=PALE, align='c')
breather(5, 2)

# узел 6
divider(6)
s = light('6.1', 'white', title='Кривая R: цена растёт с числом чтений')
text(s, 0.47, 1.15, 12.38, 0.6, '200k NDJSON · parse → normalize → consume (R проходов чтения) → serialize · процесс на ячейку · прогрев 300 / замер 1000 / 5 повторов · чексум 90/90', size=12, color=GREY)
row(s, 'white', [['R = 1', '**+26–30%**'], ['R = 5', '**×1.8**'], ['R = 20', '**×2.7–3.2**']], 1.85, 1.9, size=28, align='c')
todo(s, 0.47, 4.0, 8.4, 2.5, 'кривая по R (обе пары батчей сливаются), вторая подпись — доля parse 66% → 11%', 'ГРАФИК')
text(s, 9.15, 4.0, 3.7, 2.5, ['батч 500 = 5000', 'serialize ×1.27–1.42 при меньших байтах', 'пара той же формы 1 vs 8 — другой код, чей — в следующей главе'], size=13, color=GREY, gap=6, anchor='m')

s = light('6.2', 'green', title='HTTP: 914 против 530')
ab_cards(s, 1.3, 2.7, ['**914 rps**'], ['**530 rps**'], size=60, align='c', labels=('mono', 'guard'))
row(s, 'blue', ['p50 · 52 → 92 мс', 'p99 · 83 → 139 мс (**+67%**)'], 4.25, 0.95, x0=0.47, x1=8.87, size=18, align='c')
todo(s, 0.47, 5.4, 8.4, 1.1, 'R7 · bench-http.sh, итоговые таблицы autocannon (Req/Sec, p99)', 'ЗАПИСЬ')
text(s, 9.15, 4.25, 3.7, 2.25, ['**×1.72**', 'node:http', '`autocannon -c 50 -d 20`', 'batch 1000 · R=5', 'термин: **p99**'], size=14, color=GREY, gap=4, anchor='m')

s = light('6.3', 'grey', title='Где это не важно')
row(s, 'shadow', [['**ingest без чтений**', 'R≈0 — ~4%'], ['**I/O-bound**', 'ждёт базу десятки мс — не видно ни в p50, ни в p99'], ['**малые батчи, R=1**', '0.38 vs 0.50 мс на 500 записей']], 1.35, 2.7, size=18, gap=6)
card(s, 'dark', 0.47, 4.3, 12.38, 1.2, 'читаете нормализованное — от четверти до трёх раз; только складываете — единицы процентов', size=20, color=WHITE, align='c')
text(s, 0.47, 5.65, 12.38, 0.8, '«упрётся в базу» — верно при R≈0 · мой consume плотный — калибруйте R на свой профиль', size=14, color=GREY, align='c', anchor='m')
# замечание Влада
s = prs.slides.add_slide(LAY['white'])
set_title(s, '[резерв] Как мерили — и где d8 врёт')
drop_empty_placeholders(s)
bullets(s, 0.6, 1.4, 12.0, ['чем d8 / микробенч отличается от прода — и почему числа всё равно переносятся', 'как собирались метрики: процесс на ячейку, прогрев, медианы, чексуммы', 'ловушки замера V8: %-интринсики, --no-turbofan, общий процесс, фон харнесса'], size=20, step=1.2)
todo(s, 0.6, 5.1, 12.0, 1.4, 'слайд из обсуждения с Владом (6.09) — в сценарии 06 его нет; решить: сюда, в Q&A-резерв или в PITFALLS.md', 'РЕШИТЬ')
s.notes_slide.notes_text_frame.text = 'Замечание Влада от 6.09: доработать момент с тестированием на d8, чем отличается от реального прода; рассказать больше про сбор метрик и ловушки замера с v8/d8. В сценарии 06 отдельного блока нет.'

# узел 7
divider(7)
s = light('7.1', 'green', title='Кто написал B?')
todo(s, 0.47, 1.3, 7.6, 5.2, 'R8 · тикет P1 → `claude -p --model opus` (свежая сессия) → код → shape-analyzer. Запись по умолчанию, live только при уверенной сети', 'ЗАПИСЬ / LIVE')
for i, t in enumerate(('тикет: normalizeUser · 5 обязательных · 3 опциональных · дефолты', '`claude -p --model opus` · свежая сессия', '4 ветки × 15 = **60** генераций', 'прибор валидирован на A/B: 0.37 / 1.58 мс')):
    card(s, 'shadow', 8.35, 1.3 + i * 1.32, 4.5, 1.2, t, size=17, pad=0.25)

s = light('7.2', 'white', title='Без правил: 30/30 MEGA — и ни одной ошибки')
clone(s, 'pinkcard', 0.47, 1.3, 12.38, 5.2)
big(s, 0.9, 1.45, 11.5, 1.7, '30/30 MEGA', size=80, color='8A1538')
row(s, 'pink', ['**8** форм', 'p50 **1.60 мс** (1.57–1.68)', '**0** битых · **0** некорректных из 60'], 3.4, 1.5, x0=0.94, x1=12.4, size=19, align='c')
big(s, 0.94, 5.1, 11.46, 1.1, '«Ревью пропустит»', size=30)

s = light('7.3', 'white', title='Две записи — одна судьба')
ab_cards(s, 1.3, 3.5, ['`...(raw.referrer && {`', '`   referrer: raw.referrer })`', '×3'], ['`if (src.referrer != null)`', '`   result.referrer = src.referrer`', '×3'], size=21, labels=('B с первого слайда', 'агент, 30/30'))
row(s, 'blue', ['спред 0/30 · Object.assign 0/30 · if-guard 30/30', '1.58 ∈ [1.57 – 1.68]'], 5.05, 1.0, size=18, align='c')
text(s, 0.47, 6.15, 12.38, 0.4, 'результат генерации с 7.1: shapes 8, MEGA', size=13, color=GREY, align='c')

s = light('7.4', 'grey', title='Чей код стоял в сервисе')
card(s, 'pinkcard', 0.47, 1.3, 12.38, 1.6, '530 rps = guard-llm = `out/p1-no-rules/candidate_1.js`', size=26, align='c')
row(s, 'shadow', [['P1 / P2', '**15/15 · 15/15 MEGA**'], ['один вендор', '**N = 15**'], ['напарник (914)', '**— дальше**']], 3.15, 1.6, size=18, align='c')
todo(s, 0.47, 5.0, 12.38, 1.5, 'карта, состояние 7: автор B — агент без правил, 30/30 · автор A — ?', 'СХЕМА')

# узел 8
divider(8)
s = dark('30/30 — без единого примера в промпте', '«Добавь ключ, только если есть значение»', ['Частота в корпусе — нет замера', 'В том же корпусе перф-советы: живые и мёртвые вперемешку'], sid='8.1', big_size=40)

s = light('8.2', 'green', title='Живое: одна строка матрицы')
clone(s, 'white', 0.47, 1.3, 12.38, 5.2)
text(s, 0.9, 1.45, 11.5, 0.5, '8 форм vs 1 на чтении', size=20, align='c')
for i, (rt, v) in enumerate((('Node 20', '×3.6'), ('Node 22', '×3.2'), ('Node 24', '×4.1'), ('d8 11.3→15.2', '×3.4'))):
    x = 0.94 + i * 2.9
    card(s, 'lime', x, 2.1, 2.7, 2.0, [f'**{v}**', rt], size=26, align='c')
row(s, 'blue', ['9 / 10 строк не зависят от рантайма', '`delete` → словарный режим ×11–×32'], 4.35, 1.1, x0=0.94, x1=12.4, size=18, align='c')
text(s, 0.94, 5.6, 11.46, 0.7, '300 прогонов · процесс на кейс · таблица матрицы целиком — только фоном', size=13, color=GREY, align='c', anchor='m')

s = light('8.3', 'white', title='Мёртвое — и стык с объектами')
clone(s, 'shadow', 0.47, 1.3, 6.0, 3.3)
text(s, 0.8, 1.5, 5.4, 0.7, '«try/catch дорогой»', size=26, font='X5 Sans Medium')
text(s, 0.8, 2.4, 5.4, 2.0, ['наивный бенч: **×2.3–2.55**', 'изолированные процессы: **65 = 65 мс**'], size=20, gap=8)
todo(s, 6.75, 1.3, 6.1, 3.3, 'два прогона try/catch из эксп. 3: наивный в одном процессе / изолированный', 'СКРИН')
card(s, 'dark', 0.47, 4.85, 12.38, 1.2, 'Норма «условный ключ» нарушает живое правило. Агент живое от мёртвого не отличает', size=20, color=WHITE, align='c')
text(s, 0.47, 6.12, 12.38, 0.4, '(Node 22 с 22.9 — без Maglev: PITFALLS)', size=12, color=GREY, align='c')
breather(8, 3)

# узел 9
divider(9)
s = light('9.1', 'grey', title='Контракт: четыре правила', title_w=7)
text(s, 7.6, 0.45, 3.7, 0.4, '`v8-rules.md` — полстраницы, без кода', size=13, color=GREY, align='r')
rules = [['Один набор полей в одном порядке; нет значения → `null`'], ['После создания не добавлять и не удалять', '(`delete` ×11–×32)'], ['Однородные массивы'], ['Стабильные типы аргументов и возвратов']]
for i, r in enumerate(rules):
    x, y = (0.64, 6.89)[i % 2], (1.3, 3.55)[i // 2]
    card(s, 'shadow', x, y, 5.8, 2.05, r, size=19, pad=0.45)
    clone(s, 'badge', x + 4.98, y + 1.27, 0.62, 0.62)
    text(s, x + 4.98, y + 1.27, 0.62, 0.62, str(i + 1), size=18, align='c', anchor='m', font='X5 Sans Medium')
card(s, 'blue', 0.64, 5.8, 12.05, 0.75, 'порядок тоже деньги: 2 формы · чтение ×1.32–1.47 · +16–21% (R=20) · 849 rps (−7%)', size=15, align='c', pad=0.15)

s = light('9.2', 'green', title='Тот же агент + файл правил')
todo(s, 0.47, 1.3, 6.2, 5.2, 'LIVE №2 · генерация с v8-rules.md инлайном → литерал с null → shapes 1, MONO. Ожидание > 60 с → запись R9', 'LIVE-ДЕМО')
clone(s, 'white', 6.95, 1.3, 5.9, 5.2)
big(s, 6.95, 1.4, 5.9, 1.5, '30/30 MONO', size=54)
text(s, 7.3, 2.95, 5.2, 2.2, ['1 форма · ни одной POLY', 'p50 **0.36** (P1) / **0.38** (P2) · 0.32–0.47', '**×4.4** к 1.60 — диапазоны не пересекаются'], size=18, gap=8, align='c')
card(s, 'lime', 7.3, 5.2, 5.2, 1.05, '«отвечала за 1.6 мс — отвечает за 0.36»', size=17, align='c', pad=0.15)

s = light('9.3', 'white', title='Идиома с правилами — литерал с null')
ab_cards(s, 1.3, 3.7, ['`referrer: raw.referrer ?? null`', '×3'], ['`referrer: src.referrer != null`', '`   ? src.referrer : null`', '×3'], size=21, labels=('A с первого слайда', 'агент + правила, 30/30'))
row(s, 'lime', ['единый литерал', 'ни одного if после литерала', '**1 форма**'], 5.25, 1.1, size=19, align='c')

s = light('9.4', 'green', title='Один агент. Разница — файл')
ab_cards(s, 1.3, 2.9, ['**914 rps**', '`out/p1-rules/candidate_1.js`'], ['**530 rps**', '`out/p1-no-rules/candidate_1.js`'], size=36, align='c', labels=('с правилами · mono-llm', 'без · guard-llm'))
card(s, 'blue', 0.47, 4.45, 12.38, 0.85, '«генерации, не пара с первого слайда; одна пара форм, два прибора»', size=17, align='c', pad=0.15)
row(s, 'shadow', ['оговорка: правила инлайнились явно', 'носитель: `CLAUDE.md` · cursor rules · codex'], 5.5, 1.0, size=16, align='c', pad=0.2)

# узел 10
divider(10)
s = light('10.1', 'grey', title='Turbolev: из чего собран и в каком статусе')
card(s, 'lime', 0.47, 1.4, 4.6, 2.3, ['**Maglev**', 'фронтенд: тот же граф, те же записи → та же ставка'], size=18, align='c')
clone(s, 'chev', 5.35, 2.2, 0.26, 0.77)
card(s, 'lime', 5.9, 1.4, 4.6, 2.3, ['**Turboshaft**', 'бэкенд — напрямую'], size=18, align='c')
card(s, 'shadow', 10.8, 1.4, 2.05, 2.3, ['мимо IR', 'TurboFan'], size=16, align='c', pad=0.15)
card(s, 'dark', 0.47, 4.0, 12.38, 2.4, ['`--turbolev` = false · все ветки до 15.5', 'Chrome: A/B-эксперимент V8Turbolev · срок дефолта не объявлен', '', 'это замена верхнего яруса, не пятый ярус'], size=19, color=WHITE, align='c')

s = light('10.2', 'white', title='Числа Turbolev — только запись')
todo(s, 0.47, 1.3, 7.6, 5.2, 'R10 · d8 15.2: без флага 36 мс, с --turbolev 41.6 мс; 2 секунды в кадре %PrepareFunctionForOptimization. Никогда не live', 'ЗАПИСЬ')
for i, (t, key) in enumerate((('d8 14.9 · **×2.7** позади', 'pink'), ('d8 15.2 · **41.6 vs 36 мс**', 'lime'), ('компиляция **÷2** · о качестве кода не заявлено', 'blue'), ('A/B под `--turbolev` — **нет замера**', 'shadow'))):
    card(s, key, 8.35, 1.3 + i * 1.32, 4.5, 1.2, t, size=18, pad=0.25)

s = light('10.3', 'green', title='Ставка та же')
tl = (('Crankshaft', '2010'), ('TurboFan', '2017'), ('Sparkplug', '2021'), ('Maglev', '2023'), ('Turboshaft', '2022→23'), ('Turbolev', 'за флагом'))
for i, (n_, y_) in enumerate(tl):
    card(s, 'white' if i < 5 else 'lime', 0.47 + i * 2.08, 1.5, 1.93, 1.9, [f'**{y_}**', n_], size=17, align='c', pad=0.1)
clone(s, 'dark', 0.47, 3.75, 12.38, 2.7)
big(s, 0.47, 3.85, 3.6, 2.5, '2008', size=72, color=WHITE)
text(s, 4.2, 3.85, 8.3, 2.5, ['формы и записи мест чтения', 'старше всех ступеней лестницы'], size=22, color=PALE, anchor='m', gap=6)

# узел 11
divider(11)
s = light('11.1', 'white', title='Карта целиком: две развилки, один файл')
todo(s, 0.47, 1.3, 8.6, 5.2, 'lifecycle-map · состояние 11, полная: обе полосы закрашены, две развилки', 'СХЕМА')
for i, t in enumerate(('справа развилку задаёт **форма при сборке**', 'слева — **файл правил** в контексте агента', 'человек: файл — до агента · ревью — после', 'билдер: tsc / esbuild — нет замера, эксп. 4')):
    card(s, 'shadow' if i < 3 else 'blue', 9.3, 1.3 + i * 1.32, 3.55, 1.2, t, size=15, pad=0.22)

s = light('11.2', 'grey', title='Вывод: что агент, что человек')
row(s, 'shadow', [['**0 некорректных из 60**', 'код агента правильный'], ['**30/30**', 'идиоматика корпуса — с правилами и без'], ['**перф-сознательность**', 'остаётся человеку']], 1.5, 3.2, size=22, align='c', gap=10)
card(s, 'dark', 0.47, 5.0, 12.38, 1.4, 'файл правил — до агента · взгляд на набор ключей — на ревью', size=22, color=WHITE, align='c')

s = light('11.3', 'green', title='Как проверить у себя')
for i, (cmd, what) in enumerate((('`node --allow-natives-syntax lab/shape-analyzer.js <файл>`', '%HaveSameMap → число форм'), ('`node --log-ic --no-logfile-per-isolate --logfile=ic.log …`', 'состояние мест чтения: MONO / MEGA'), ('`bash run-matrix.sh`', 'матрицу → на своём Node'))):
    clone(s, 'white', 0.47, 1.3 + i * 1.45, 12.38, 1.3)
    text(s, 0.85, 1.3 + i * 1.45, 7.6, 1.3, cmd, size=16, anchor='m')
    text(s, 8.7, 1.3 + i * 1.45, 3.9, 1.3, what, size=17, anchor='m', font='X5 Sans Medium')
text(s, 0.47, 5.75, 12.38, 0.7, 'Q&A-резерв: compile hints · окно Maglev · ×27 · HOLEY · delete на d8 15.2 · две ×4.4', size=13, color=GREY, align='c', anchor='m')

# узел 12
divider(12)
s = light('12.1', 'white', title='Две функции, две судьбы')
ab_cards(s, 1.3, 1.9, ['`referrer: raw.referrer ?? null`', '**0.37 мс**'], ['`...(raw.referrer && { referrer: … })`', '**1.58 мс**'], size=20, align='c')
todo(s, 0.47, 3.4, 12.38, 2.5, 'lifecycle-map · полная карта (состояние 11) рядом с первым слайдом', 'СХЕМА')
text(s, 0.47, 6.0, 12.38, 0.6, '0.37 / 1.58 — функция (shape-analyzer) · 914 / 530 — сервис (pipeline + autocannon) · одна пара форм 1 vs 8', size=14, color=GREY, align='c', anchor='m')

s = dark('ветка v8-holy-autumn · PITFALLS.md', 'Спасибо', ['github.com/Yoshimuro/issues-memory-research', 'Александр Зайцев'], sid='12.2')
todo(s, 10.2, 1.0, 2.4, 2.4, 'QR на ветку репо', 'QR')

# Q&A-резерв
s = prs.slides.add_slide(LAY['grey'])
set_title(s, 'Q&A-резерв')
drop_empty_placeholders(s)
qa = ['compile hints — не про ярусы', 'окно Maglev ~260..10 000 — гипотеза', '×27', 'HOLEY-массивы', '`delete` на d8 15.2', '«две ×4.4»', 'Sparkplug не спекулирует', 'сколько деоптов? — счётчик забракован']
for i, t in enumerate(qa):
    card(s, 'shadow', 0.47 + (i % 4) * 3.13, 1.4 + (i // 4) * 2.55, 2.98, 2.35, t, size=18, align='c', pad=0.2)
s.notes_slide.notes_text_frame.text = 'Заготовки из 06-talk-final.md: «Узел 11 → Q&A-резерв» и «Риски сцены → Сеньор в Q&A». Каждую плашку развернуть в отдельный слайд по мере готовности.'

# ---------- выкинуть слайды шаблона ----------
lst = prs.slides._sldIdLst
for sldId in list(lst)[:N_TPL]:
    prs.part.drop_rel(sldId.rId)
    lst.remove(sldId)
prs.save(OUT)
print('slides:', len(prs.slides), '| parsed:', len(SL), 'slides,', len(NODES), 'nodes')
